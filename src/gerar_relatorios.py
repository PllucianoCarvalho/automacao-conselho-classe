import argparse
import logging
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt
from tqdm.auto import tqdm


CONFIG = {
    "abas_turmas": ["3A", "6A", "6B", "6C", "7A", "7B", "7C", "8A", "8B", "8C", "9A"],
    "colunas_esperadas": [
        "Nº",
        "ALUNO",
        "NIVEL DE PROEFICIENCIA DO ALUNO",
        "COMPROMETIDO COM APRENDIZAGEM?",
        "PROATIVO E/OU CURIOSO?",
        "PARTICIPATIVO?",
        "EDUCADO(A)?",
        "FALTOSO?",
        "POSSUI RELATORIO?",
        "EM QUAL ÁREA APRESENTA HABILIDADE?",
        "NIVEL DO ALUNO SEGUNDO AS RESPOSTAS",
    ],
    "colunas_numericas": [
        "NIVEL DE PROEFICIENCIA DO ALUNO",
        "COMPROMETIDO COM APRENDIZAGEM?",
        "PROATIVO E/OU CURIOSO?",
        "PARTICIPATIVO?",
        "EDUCADO(A)?",
        "FALTOSO?",
    ],
}

DESCRICOES = {
    "NIVEL DE PROEFICIENCIA DO ALUNO": [
        "sem dados sobre nível de proficiência",
        "tem nível abaixo do básico em aprendizagem",
        "tem nível básico em aprendizagem",
        "tem nível adequado em aprendizagem",
        "tem nível avançado em aprendizagem",
    ],
    "COMPROMETIDO COM APRENDIZAGEM?": [
        "sem dados sobre comprometimento",
        "com baixo comprometimento",
        "com comprometimento básico",
        "com comprometimento adequado",
        "com alto comprometimento",
    ],
    "PROATIVO E/OU CURIOSO?": [
        "sem dados sobre proatividade/curiosidade",
        "com baixa proatividade ou curiosidade",
        "com proatividade ou curiosidade moderada",
        "com proatividade ou curiosidade adequadas",
        "com alta proatividade e curiosidade",
    ],
    "PARTICIPATIVO?": [
        "sem dados sobre participação",
        "com baixa participação nas atividades",
        "com participação moderada nas atividades",
        "com participação adequada nas atividades",
        "com alta participação nas atividades",
    ],
    "EDUCADO(A)?": [
        "sem dados sobre comportamento",
        "com comportamento inadequado",
        "com comportamento básico",
        "com comportamento adequado",
        "com comportamento exemplar e educado",
    ],
    "FALTOSO?": [
        "sem dados sobre presença",
        "com excelente assiduidade",
        "com boa assiduidade e poucas faltas",
        "com assiduidade moderada e algumas faltas",
        "com baixa assiduidade e muitas faltas",
    ],
}

COLUNA_HABILIDADES = "EM QUAL ÁREA APRESENTA HABILIDADE?"


def descrever_criterio(valor, criterio):
    if pd.isna(valor) or valor == "":
        return DESCRICOES[criterio][0] + "; "
    try:
        nivel = round(float(valor))
    except (ValueError, TypeError):
        return f"{criterio.lower()} inválido; "
    if 0 <= nivel <= 4:
        return DESCRICOES[criterio][nivel] + "; "
    return f"{criterio.lower()} inválido; "


def gerar_relatorio(row):
    nome = str(row.get("ALUNO", "")).strip()
    if not nome:
        return "N/A"

    descricoes = {
        criterio: descrever_criterio(row.get(criterio), criterio)
        for criterio in CONFIG["colunas_numericas"]
    }
    criterios_validos = sum(
        "sem dados" not in descricao.lower() and "inválido" not in descricao.lower()
        for descricao in descricoes.values()
    )

    educacao = row.get("EDUCADO(A)?")
    transicao_educacao = ""
    if pd.notna(educacao) and educacao >= 3:
        transicao_educacao = "Por outro lado, "

    faltas = row.get("FALTOSO?")
    transicao_faltas = ""
    if pd.notna(faltas) and faltas <= 2:
        transicao_faltas = "No entanto, "

    habilidade = str(row.get("HABILIDADES", "N/A")).strip()
    relatorio_habilidades = (
        f" Demonstra habilidade em {habilidade.lower()}."
        if habilidade and habilidade.upper() != "N/A"
        else ""
    )

    possui_relatorio = str(row.get("POSSUI RELATORIO?", "")).strip().upper()
    if possui_relatorio == "SIM":
        relatorio_relatorio = " Aluno com relatório."
    elif possui_relatorio == "NÃO":
        relatorio_relatorio = ""
    else:
        relatorio_relatorio = " Sem informação sobre relatório."

    relatorio = (
        f"{nome} {descricoes[CONFIG['colunas_numericas'][0]]}"
        f"{descricoes[CONFIG['colunas_numericas'][1]]}"
        f"{descricoes[CONFIG['colunas_numericas'][2]]}"
        f"{descricoes[CONFIG['colunas_numericas'][3]]}"
        f"{transicao_educacao}{descricoes[CONFIG['colunas_numericas'][4]]}"
        f"{transicao_faltas}{descricoes[CONFIG['colunas_numericas'][5]]}"
        f"{relatorio_relatorio}{relatorio_habilidades}"
    ).strip().replace("; ", ". ").replace("..", ".").replace(" .", ".").capitalize()

    if criterios_validos == 0 and possui_relatorio != "SIM" and habilidade.upper() == "N/A":
        return f"{nome.title()}. Dados insuficientes para gerar relatório."
    return relatorio


def padronizar_nome(nome):
    if pd.isna(nome) or nome == "":
        return ""
    return str(nome).strip().upper()


def validar_valores_numericos(df, colunas_numericas):
    for coluna in colunas_numericas:
        df[coluna] = pd.to_numeric(df[coluna], errors="coerce")
        invalidos = df[coluna].notna() & ~df[coluna].between(0, 4)
        if invalidos.any():
            logging.warning(
                "Valores fora do intervalo na coluna %s: %s",
                coluna,
                df.loc[invalidos, coluna].tolist(),
            )
            df.loc[invalidos, coluna] = pd.NA
    return df


def limpar_dados(df, aba, nome_arquivo):
    tamanho_inicial = len(df)
    df = df.dropna(subset=["ALUNO"])
    df = df[df["ALUNO"].str.strip() != ""]
    if len(df) < tamanho_inicial:
        logging.info(
            "%s linhas removidas da aba %s do arquivo %s por ALUNO vazio.",
            tamanho_inicial - len(df),
            aba,
            nome_arquivo,
        )
    return df


def combinar_habilidades(*habilidades):
    habilidades_set = set()
    for habilidade in habilidades:
        if isinstance(habilidade, str) and habilidade.strip() and habilidade.upper() != "N/A":
            habilidades_set.update(
                item.strip().upper()
                for item in habilidade.split(",")
                if item.strip()
            )
    return ", ".join(sorted(habilidades_set)) if habilidades_set else "N/A"


def combinar_dfs_por_aba(dfs, aba):
    if not dfs:
        logging.warning("Nenhum dado válido para a aba %s.", aba)
        return None

    dataframes_validos = []
    for nome_arquivo, df in dfs:
        df = df.copy()
        df["ALUNO"] = df["ALUNO"].apply(padronizar_nome)
        df = validar_valores_numericos(df, CONFIG["colunas_numericas"])
        df = limpar_dados(df, aba, nome_arquivo)
        if not df.empty:
            dataframes_validos.append(df)

    if not dataframes_validos:
        logging.warning("Todos os dados da aba %s ficaram vazios após a limpeza.", aba)
        return None

    df_concat = pd.concat(dataframes_validos, ignore_index=True)
    agregacoes = {
        coluna: "mean"
        for coluna in CONFIG["colunas_numericas"]
    }
    agregacoes.update(
        {
            "Nº": "first",
            "POSSUI RELATORIO?": lambda valores: (
                "SIM"
                if any(str(valor).strip().upper() == "SIM" for valor in valores if pd.notna(valor))
                else "NÃO"
            ),
            COLUNA_HABILIDADES: lambda valores: combinar_habilidades(
                *[valor for valor in valores if pd.notna(valor)]
            ),
        }
    )

    try:
        df_merged = df_concat.groupby("ALUNO", as_index=False).agg(agregacoes)
    except Exception:
        logging.exception("Erro ao combinar dados para a aba %s.", aba)
        return None

    tem_dados_validos = any(
        df_merged[coluna].notna().any()
        for coluna in CONFIG["colunas_numericas"]
    ) or df_merged["POSSUI RELATORIO?"].eq("SIM").any() or df_merged[COLUNA_HABILIDADES].ne("N/A").any()
    if not tem_dados_validos:
        logging.warning("Nenhum indicador válido encontrado para a aba %s.", aba)
        return None

    logging.info("Dados combinados para a aba %s: %s alunos.", aba, len(df_merged))
    return df_merged


def verificar_abas(caminho_arquivo, abas_turmas):
    try:
        with pd.ExcelFile(caminho_arquivo) as planilha:
            abas_encontradas = [aba for aba in planilha.sheet_names if aba in abas_turmas]
        logging.info("Abas encontradas em %s: %s", caminho_arquivo.name, abas_encontradas)
        return abas_encontradas
    except Exception:
        logging.exception("Erro ao verificar abas no arquivo %s.", caminho_arquivo)
        return []


def validar_dataframe(df, aba, nome_arquivo):
    if "ALUNO" not in df.columns:
        logging.error("Coluna ALUNO ausente na aba %s do arquivo %s.", aba, nome_arquivo)
        return False
    if df["ALUNO"].dropna().empty:
        logging.error("Coluna ALUNO vazia na aba %s do arquivo %s.", aba, nome_arquivo)
        return False

    colunas_ausentes = [coluna for coluna in CONFIG["colunas_esperadas"] if coluna not in df.columns]
    if colunas_ausentes:
        logging.warning(
            "Colunas ausentes na aba %s do arquivo %s: %s. Serão consideradas vazias.",
            aba,
            nome_arquivo,
            colunas_ausentes,
        )
    for coluna in colunas_ausentes:
        df[coluna] = pd.NA
    return True


def ler_aba(caminho_arquivo, aba):
    df = pd.read_excel(caminho_arquivo, sheet_name=aba, header=None, skiprows=3)
    colunas = CONFIG["colunas_esperadas"] + [
        f"Unnamed_{indice}"
        for indice in range(len(CONFIG["colunas_esperadas"]), len(df.columns))
    ]
    df.columns = colunas[: len(df.columns)]
    return df


def criar_documento(df_resultado, aba, pasta_saida):
    documento = Document()
    documento.add_heading(f"Relatórios dos Alunos - Turma {aba}", level=1)
    tabela = documento.add_table(rows=1, cols=3)
    tabela.style = "Table Grid"

    for celula, coluna in zip(tabela.rows[0].cells, ["Nº", "ALUNO", "RELATÓRIO"]):
        celula.text = coluna
        for run in celula.paragraphs[0].runs:
            run.font.size = Pt(12)
            run.bold = True

    for _, row in df_resultado.iterrows():
        celulas = tabela.add_row().cells
        celulas[0].text = str(row.get("Nº", "N/A"))
        celulas[1].text = str(row.get("ALUNO", "N/A")).title()
        celulas[2].text = str(row.get("RELATÓRIO", "N/A"))
        for celula in celulas:
            for paragrafo in celula.paragraphs:
                for run in paragrafo.runs:
                    run.font.size = Pt(10)

    caminho_saida = pasta_saida / f"relatorios_{aba}.docx"
    documento.save(caminho_saida)
    logging.info("Arquivo %s criado com %s alunos.", caminho_saida, len(df_resultado))
    return caminho_saida


def processar_arquivos(pasta_entrada, pasta_saida):
    arquivos = sorted(
        caminho
        for caminho in pasta_entrada.iterdir()
        if caminho.is_file() and caminho.suffix.lower() == ".xlsx"
    )
    if not arquivos:
        raise FileNotFoundError(f"Nenhum arquivo .xlsx encontrado em {pasta_entrada}.")

    pasta_saida.mkdir(parents=True, exist_ok=True)
    dfs_por_aba = {aba: [] for aba in CONFIG["abas_turmas"]}

    for caminho_arquivo in tqdm(arquivos, desc="Lendo arquivos Excel"):
        abas_presentes = verificar_abas(caminho_arquivo, CONFIG["abas_turmas"])
        for aba in abas_presentes:
            try:
                df = ler_aba(caminho_arquivo, aba)
                if validar_dataframe(df, aba, caminho_arquivo.name):
                    dfs_por_aba[aba].append((caminho_arquivo.name, df))
            except Exception:
                logging.exception("Erro ao ler a aba %s do arquivo %s.", aba, caminho_arquivo.name)

    relatorios_criados = []
    for aba in tqdm(CONFIG["abas_turmas"], desc="Processando turmas"):
        df_merged = combinar_dfs_por_aba(dfs_por_aba[aba], aba)
        if df_merged is None or df_merged.empty:
            print(f"Turma {aba}: nenhum dado válido encontrado.")
            continue

        df_resultado = pd.DataFrame(
            {
                "Nº": df_merged["Nº"].fillna("SEM NÚMERO"),
                "ALUNO": df_merged["ALUNO"].str.title(),
                "POSSUI RELATORIO?": df_merged["POSSUI RELATORIO?"],
                "HABILIDADES": df_merged[COLUNA_HABILIDADES],
            }
        )
        for coluna in CONFIG["colunas_numericas"]:
            df_resultado[coluna] = df_merged[coluna]
        df_resultado["RELATÓRIO"] = df_resultado.apply(gerar_relatorio, axis=1)
        relatorios_criados.append(criar_documento(df_resultado, aba, pasta_saida))
        print(f"Turma {aba}: {len(df_resultado)} alunos processados.")

    return relatorios_criados


def main():
    parser = argparse.ArgumentParser(
        description="Consolida planilhas de conselhos de classe e gera relatórios DOCX por turma."
    )
    parser.add_argument("--entrada", type=Path, default=Path("dados"), help="Pasta com arquivos .xlsx.")
    parser.add_argument(
        "--saida",
        type=Path,
        default=Path("relatorios_gerados"),
        help="Pasta onde os relatórios .docx serão salvos.",
    )
    args = parser.parse_args()

    logging.basicConfig(
        filename="processamento_relatorio.log",
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
    )

    if not args.entrada.is_dir():
        parser.error(f"A pasta de entrada não existe: {args.entrada}")

    try:
        relatorios = processar_arquivos(args.entrada, args.saida)
    except FileNotFoundError as erro:
        parser.error(str(erro))

    print(f"\nProcessamento concluído. {len(relatorios)} relatório(s) salvo(s) em {args.saida}.")


if __name__ == "__main__":
    main()